"""Component characteristic plots for the engineering-notes tab.
Combines the team's earlier spreadsheet data with the automated run of 10 Oct 2026."""
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import openpyxl

HERE = Path(__file__).resolve().parent
NOTES = Path(r"C:\Users\Sorin\Downloads\Engineering notes-20261010T221011Z-1-001\Engineering notes\Experimental Data")
OUT = HERE / "site" / "img"
BAND = (1000, 1200)  # MHz, the IF band the design wants

run3 = np.genfromtxt(HERE / "phase_detector_calibration_run3_cleanDC_0p02V_steps.csv", delimiter=",", names=True, dtype=None, encoding="utf-8")
run3 = run3[run3["direction"] == "up"]


def col(ws, c, r0, r1):
    return [ws[f"{c}{r}"].value for r in range(r0, r1 + 1)]


# ---- VCO tuning curve: spreadsheet (0-10 V, 4 Oct) + automated run (0-3.5 V, 10 Oct)
ws = openpyxl.load_workbook(NOTES / "VCO tests.xlsx", data_only=True)["absolute accuracy"]
v_team = np.array(col(ws, "B", 3, 22), float)
f_team = np.array(col(ws, "D", 3, 22), float) * 1e3
fig, ax = plt.subplots(figsize=(8.5, 5.2))
ax.axhspan(*BAND, color="tab:orange", alpha=0.15, label="band the design wants, 1.0-1.2 GHz")
ax.plot(v_team, f_team, "s-", color="tab:gray", ms=5, label="4 Oct, by hand, VCO straight into the analyzer")
ax.plot(run3["vt_meas_v"], run3["vco_freq_hz"] / 1e6, "-", color="tab:green", lw=2.2, label="10 Oct, automated, VCO into tee + splitter + mixer")
ax.plot(v_team, 800 + 60 * v_team, ":", color="k", lw=1, label="straight line 0.8 GHz + 60 MHz/V")
lo, hi = np.interp(BAND, f_team, v_team)
ax.axvspan(lo, hi, color="tab:orange", alpha=0.08)
ax.annotate(f"needs {lo:.1f} to {hi:.1f} V", (0.5 * (lo + hi), BAND[0] - 55), ha="center", fontsize=9, color="tab:orange")
ax.set(xlabel="Tuning voltage Vt (V)", ylabel="Output frequency (MHz)", title="VCO (Crystek CVCO55BE-0800-1600): frequency vs tuning voltage")
ax.grid(alpha=0.3)
ax.legend(fontsize=8)
fig.tight_layout()
fig.savefig(OUT / "comp_vco_tuning.png", dpi=125)
print("band needs %.2f to %.2f V; slope in band %.1f MHz/V" % (lo, hi, (BAND[1] - BAND[0]) / (hi - lo)))
d = f_team - (800 + 60 * v_team)
print("straight-line error: %.0f to %.0f MHz" % (d.min(), d.max()))
m = v_team <= 3.5
print("4 Oct minus 10 Oct at same voltage: %s MHz" % np.round(f_team[m] - np.interp(v_team[m], run3["vt_meas_v"], run3["vco_freq_hz"] / 1e6), 1))

# ---- VCO output power vs frequency (team, 6 Oct)
ws = openpyxl.load_workbook(NOTES / "Attenuation tracker.xlsx", data_only=True)["VCO Power"]
fp = np.array(col(ws, "E", 3, 12), float) * 1e3
pw = np.array(col(ws, "F", 3, 12), float)
o = np.argsort(fp)
fig, ax = plt.subplots(figsize=(8.5, 4.6))
ax.axvspan(*BAND, color="tab:orange", alpha=0.15, label="1.0-1.2 GHz")
ax.plot(fp[o], pw[o], "o-", color="tab:blue", label="measured 6 Oct, VCO + cable into the analyzer")
ax.axhspan(4, 8, color="tab:green", alpha=0.08, label="listed output, 6 +/- 2 dBm")
ax.set(xlabel="Frequency (MHz)", ylabel="Output power (dBm)", title="VCO output power vs frequency")
sec = ax.secondary_yaxis("right", functions=(lambda x: 10 ** (x / 10), lambda y: 10 * np.log10(np.maximum(y, 1e-9))))
sec.set_ylabel("mW")
ax.grid(alpha=0.3)
ax.legend(fontsize=8)
fig.tight_layout()
fig.savefig(OUT / "comp_vco_power.png", dpi=125)
inb = (fp >= BAND[0]) & (fp <= BAND[1])
print("VCO power in band: %.2f to %.2f dBm (%.1f to %.1f mW)" % (pw[inb].min(), pw[inb].max(), 10 ** (pw[inb].min() / 10), 10 ** (pw[inb].max() / 10)))

# ---- phase detector: Pluto-driven curves at four powers (team, Sept) + VCO-driven curve (10 Oct)
ws = openpyxl.load_workbook(NOTES / "Phase Detector Test.xlsx", data_only=True)["data"]
fig, ax = plt.subplots(1, 2, figsize=(15, 5.2))
for (fc, vc, label), color in zip((("B", "C", "-2 dBm"), ("F", "G", "-6 dBm"), ("K", "L", "-11 dBm"), ("P", "Q", "-12.25 dBm")),
                                  ("tab:red", "tab:orange", "tab:purple", "tab:blue")):
    pts = [(f, v) for f, v in zip(col(ws, fc, 6, 64), col(ws, vc, 6, 64)) if isinstance(f, (int, float)) and isinstance(v, (int, float))]
    f_, v_ = np.array(pts).T
    ax[0].plot(f_ * 1e3, v_, "o-", ms=3.5, color=color, label=f"Pluto set to {label}")
    print(label, "range %.0f-%.0f MHz, %.1f-%.1f mV" % (f_.min() * 1e3, f_.max() * 1e3, v_.min(), v_.max()))
ax[0].axvspan(*BAND, color="tab:orange", alpha=0.12)
ax[0].set(xlabel="Input frequency (MHz)", ylabel="Mixer output (mV, multimeter)",
          title="September: Pluto SDR as the source, four power settings")
ax[0].legend(fontsize=8)
ax[1].plot(run3["vco_freq_hz"] / 1e6, run3["mixer_out_v"] * 1e3, "-", color="tab:green", lw=2, label="VCO as the source (about +0.6 dBm at each mixer port)")
ax[1].axvspan(743.6, 788.4, color="tab:green", alpha=0.12, label="single-valued: 744-788 MHz")
ax[1].set(xlabel="Input frequency (MHz)", ylabel="Mixer output (mV, scope)", title="10 October: VCO as the source")
ax[1].legend(fontsize=8)
for a in ax:
    a.grid(alpha=0.3)
fig.suptitle("Phase detector (ZFRSC-42-S+ splitter, 4.5 in extra cable, ZX05-30W-S+ mixer): output voltage vs input frequency", fontweight="bold")
fig.tight_layout()
fig.savefig(OUT / "comp_phase_detector.png", dpi=125)

# ---- path attenuation table (team, 6 Oct)
ws = openpyxl.load_workbook(NOTES / "Attenuation tracker.xlsx", data_only=True)["Path attenuation"]
for r in range(3, 9):
    print([ws[f"{c}{r}"].value for c in "BCD"])
