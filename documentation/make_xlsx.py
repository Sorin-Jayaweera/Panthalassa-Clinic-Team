"""Build JACON_component_specs.xlsx from specs.py (same source as the Word doc)."""
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, ".")
import specs
from specs import PARTS, damage_dbm_mw

FONT = "Arial"
HDR_FILL = PatternFill("solid", fgColor="1F5FBF")
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")
BAD_FILL = PatternFill("solid", fgColor="F8CBAD")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

wb = Workbook()


def style_header(ws, row=1):
    for cell in ws[row]:
        if cell.value is None:
            continue
        cell.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        cell.fill = HDR_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[row].height = 32
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def body_style(ws, first_row=2):
    for row in ws.iter_rows(min_row=first_row):
        for cell in row:
            cell.font = Font(name=FONT, size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = BORDER


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def link(ws, cell, url, text):
    c = ws[cell]
    c.value = text
    c.hyperlink = url
    c.font = Font(name=FONT, size=10, color="1F5FBF", underline="single")


# ------------------------------------------------------------------ Components
ws = wb.active
ws.title = "Components"
HEADERS = ["Part number", "Role", "JACON block", "Datasheet band", "Frequency seen in JACON",
           "Gain / loss at that frequency", "Damage limit (as printed on datasheet)",
           "Damage limit (dBm)", "Damage limit (mW)", "P1dB / compression (dBm)", "P1dB (mW)",
           "Notes", "Datasheet page", "Mini-Circuits PDF", "Local PDF copy", "Product page",
           "Evidence screenshot"]
ws.append(HEADERS)

for pt in PARTS:
    d_dbm, _ = damage_dbm_mw(pt)
    r = ws.max_row + 1
    ws.cell(row=r, column=1, value=pt["model"])
    ws.cell(row=r, column=2, value=pt["role"])
    ws.cell(row=r, column=3, value=pt["jacon"])
    ws.cell(row=r, column=4, value=pt["band"])
    ws.cell(row=r, column=5, value=pt["f_seen"])
    ws.cell(row=r, column=6, value=pt["gain_txt"])
    ws.cell(row=r, column=7, value=pt["damage_txt"])
    ws.cell(row=r, column=8, value=round(d_dbm, 2))
    # mW is computed by the sheet, so editing the dBm cell updates it
    ws.cell(row=r, column=9, value="=10^(H%d/10)" % r)
    if pt["p1db_dbm"] is None:
        ws.cell(row=r, column=10, value="passive - n/a")
        ws.cell(row=r, column=11, value="n/a")
    else:
        ws.cell(row=r, column=10, value=pt["p1db_dbm"])
        ws.cell(row=r, column=11, value="=10^(J%d/10)" % r)
    ws.cell(row=r, column=12, value=pt["extra"])
    ws.cell(row=r, column=13, value="p.%d" % pt["damage_page"])
    model_url = pt["model"].replace("+", "%2B")
    link(ws, "N%d" % r, specs.PDF.format(model_url), "%s.pdf" % pt["model"])
    link(ws, "O%d" % r, "datasheets/%s.pdf" % pt["model"], "datasheets/%s.pdf" % pt["model"])
    link(ws, "P%d" % r, specs.URL.format(model_url), "minicircuits.com")
    ws.cell(row=r, column=17, value="evidence/%s" % pt["ev"][0][1])

body_style(ws)
style_header(ws)
for r in range(2, ws.max_row + 1):
    ws.cell(row=r, column=9).number_format = "0.000"
    ws.cell(row=r, column=8).number_format = "+0.00;-0.00"
    if isinstance(ws.cell(row=r, column=10).value, (int, float)):
        ws.cell(row=r, column=10).number_format = "+0.0;-0.0"
        ws.cell(row=r, column=11).number_format = "0.0"
    if ws.cell(row=r, column=1).value == "ZX60-2534MA-S+":
        for c in range(1, len(HEADERS) + 1):
            ws.cell(row=r, column=c).fill = BAD_FILL
widths(ws, {"A": 17, "B": 26, "C": 22, "D": 22, "E": 24, "F": 30, "G": 34, "H": 11, "I": 11,
            "J": 13, "K": 9, "L": 34, "M": 8, "N": 20, "O": 26, "P": 15, "Q": 34})
ws.auto_filter.ref = "A1:Q%d" % ws.max_row

r = ws.max_row + 2
ws.cell(row=r, column=1, value="Damage limit = the level that destroys the part (absolute maximum rating). "
        "P1dB = where it stops being linear, well below damage. mW columns are formulas: =10^(dBm/10).")
ws.cell(row=r, column=1).font = Font(name=FONT, size=9, italic=True)
ws.cell(row=r + 1, column=1, value="Every value was extracted from the PDF in datasheets/ at the page in column M. "
        "Orange row = lowest damage limit in the set; see the Headroom tab.")
ws.cell(row=r + 1, column=1).font = Font(name=FONT, size=9, italic=True)

# ------------------------------------------------------------------ Headroom
hs = wb.create_sheet("Headroom")
hs.append(["Part", "Expected worst-case input", "Expected level (dBm)", "Damage limit (dBm)",
           "Margin (dB)", "Verdict", "Action"])
HEAD = [
    ("ZX60-5123LN+ (LNA)", "limiter output while the limiter clamps", 16.3, 15.0,
     "Keep the in-line attenuator the drawing calls for."),
    ("ZX05-14LH-S+ (RX mixer)", "LNA saturated output at close range", 17.5, 17.0,
     "Marginal at close range; the same pad fixes this."),
    ("ZX60-P162LN+ (IF amp)", "mixer IF output, worst case", 11.0, 25.0, "Fine."),
    ("ZX60-P103LN+ (IF amp)", "mixer IF output, worst case", 11.0, 21.0, "Fine."),
    ("ZX60-2534MA-S+ (IF amp)", "mixer IF output, worst case", 11.0, -15.0,
     "Do not put this straight after the mixer."),
    ("ZFRSC-42-S+ (both units)", "IF amp output, saturated", 20.0, 28.75, "Fine."),
    ("ZX05-30W-S+ (phase detector)", "limiting-amp chain, per port", 8.0, 16.99,
     "Above the +1 dBm compression point, which is intended here."),
    ("ZX60-14012L-S+ (LO buffer)", "PLL output (assumed +3 dBm)", 3.0, 10.0,
     "Confirm your ADF5355 output level."),
    ("ZX10-2-1252-S+ (LO splitter)", "LO buffer output", 14.0, 32.67, "Fine."),
]
for part, cond, lvl, lim, action in HEAD:
    r = hs.max_row + 1
    hs.cell(row=r, column=1, value=part)
    hs.cell(row=r, column=2, value=cond)
    hs.cell(row=r, column=3, value=lvl)
    hs.cell(row=r, column=4, value=lim)
    hs.cell(row=r, column=5, value="=D%d-C%d" % (r, r))
    hs.cell(row=r, column=6, value='=IF(E%d<0,"OVER LIMIT",IF(E%d<3,"TIGHT","OK"))' % (r, r))
    hs.cell(row=r, column=7, value=action)
body_style(hs)
style_header(hs)
for r in range(2, hs.max_row + 1):
    for c in (3, 4, 5):
        hs.cell(row=r, column=c).number_format = "+0.0;-0.0"
    lim = hs.cell(row=r, column=4).value
    lvl = hs.cell(row=r, column=3).value
    if lim - lvl < 0:
        for c in range(1, 8):
            hs.cell(row=r, column=c).fill = BAD_FILL
    elif lim - lvl < 3:
        for c in range(1, 8):
            hs.cell(row=r, column=c).fill = WARN_FILL
widths(hs, {"A": 28, "B": 34, "C": 17, "D": 17, "E": 12, "F": 13, "G": 44})
r = hs.max_row + 2
hs.cell(row=r, column=1, value="Expected levels are simulated (LTspice, simulations/claude/01_jacon_system), "
        "not measured. Margin and verdict are formulas, so edit column C with your own measured levels and the "
        "verdicts update.")
hs.cell(row=r, column=1).font = Font(name=FONT, size=9, italic=True)

# ------------------------------------------------------------------ Evidence
es = wb.create_sheet("Evidence")
es.append(["Part number", "What it proves", "Datasheet page", "Screenshot file", "Source PDF (online)",
           "Local PDF copy"])
for pt in PARTS:
    for tag, fn, page, cap in pt["ev"]:
        r = es.max_row + 1
        es.cell(row=r, column=1, value=pt["model"])
        es.cell(row=r, column=2, value=cap)
        es.cell(row=r, column=3, value=page)
        es.cell(row=r, column=4, value="evidence/%s" % fn)
        model_url = pt["model"].replace("+", "%2B")
        link(es, "E%d" % r, specs.PDF.format(model_url), "%s.pdf" % pt["model"])
        link(es, "F%d" % r, "datasheets/%s.pdf" % pt["model"], "datasheets/%s.pdf" % pt["model"])
body_style(es)
style_header(es)
widths(es, {"A": 17, "B": 62, "C": 14, "D": 40, "E": 20, "F": 26})
r = es.max_row + 2
es.cell(row=r, column=1, value="Each screenshot was cropped directly from the PDF in datasheets/ at the listed "
        "page. Open the PDF at that page to check the value yourself.")
es.cell(row=r, column=1).font = Font(name=FONT, size=9, italic=True)

out = "JACON_component_specs.xlsx"
wb.save(out)
print("wrote", out, os.path.getsize(out), "bytes")
